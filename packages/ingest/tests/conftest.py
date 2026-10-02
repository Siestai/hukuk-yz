"""Synthetic fixtures only: real corpus files are never copied into the repo."""

import struct
import zipfile
from pathlib import Path

import pytest
from docx import Document

# Latin-5 letters need explicit glyph names, so the fixture font carries a /Differences table.
_DIFFS = {
    "ğ": ("gbreve", 128),
    "ı": ("dotlessi", 129),
    "İ": ("Idotaccent", 130),
    "ş": ("scedilla", 131),
    "ç": ("ccedilla", 132),
    "ö": ("odieresis", 133),
    "ü": ("udieresis", 134),
}

_TURKISH_SENTENCE = "İşçinin kıdem tazminatı için bir yıl çalışması gerekir ve bu süre ğş çöü"


def make_pdf(pages: list[str]) -> bytes:
    objs: list[bytes] = []
    n_pages = len(pages)
    first_page = 4
    kids = " ".join(f"{first_page + 2 * i} 0 R" for i in range(n_pages))
    diffs = " ".join(f"{code} /{name}" for name, code in _DIFFS.values())
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    objs.append(
        f"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding << /Type /Encoding "
        f"/BaseEncoding /WinAnsiEncoding /Differences [{diffs}] >> >>".encode()
    )
    for i, text in enumerate(pages):
        page_id = first_page + 2 * i
        content_id = page_id + 1
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        encoded = bytearray()
        for ch in text:
            encoded += bytes([_DIFFS[ch][1]]) if ch in _DIFFS else ch.encode("ascii")
        escaped = bytes(encoded).replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
        stream = b"BT /F1 12 Tf 50 700 Td (" + escaped + b") Tj ET" if text else b""
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for num, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % num + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


def make_ole2(*stream_names: str) -> bytes:
    """Minimal OLE2 compound file (v3): one FAT sector, one directory sector, empty streams."""
    assert len(stream_names) <= 3
    free = struct.pack("<I", 0xFFFFFFFF)
    header = bytearray(512)
    header[0:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<HHHHH", header, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<III", header, 44, 1, 1, 0)  # FAT sectors, first dir sector, tx sig
    struct.pack_into("<IIIII", header, 56, 4096, 0xFFFFFFFE, 0, 0xFFFFFFFE, 0)
    header[76:80] = struct.pack("<I", 0)  # DIFAT[0] -> FAT is sector 0
    header[80:512] = free * 108
    fat = struct.pack("<II", 0xFFFFFFFD, 0xFFFFFFFE) + free * 126

    def entry(name: str, kind: int) -> bytes:
        raw = name.encode("utf-16-le")
        body = bytearray(128)
        body[: len(raw)] = raw
        struct.pack_into("<HB", body, 64, len(raw) + 2, kind)
        return bytes(body)

    directory = entry("Root Entry", 5) + b"".join(entry(n, 2) for n in stream_names)
    return bytes(header) + fat + directory.ljust(512, b"\x00")


def make_zip(path: Path, members: dict[str, str | bytes]) -> Path:
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return path


@pytest.fixture
def corpus(tmp_path: Path) -> Path:
    root = tmp_path / "corpus"
    root.mkdir()
    long_text = " ".join([_TURKISH_SENTENCE] * 3)
    (root / "text.pdf").write_bytes(make_pdf([long_text, long_text]))
    (root / "scan.pdf").write_bytes(make_pdf(["", ""]))

    doc = Document()
    doc.add_paragraph("Madde 1 - İşveren ve işçi ğış")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "ücret"
    table.rows[0].cells[1].text = "ıslak"
    doc.save(str(root / "a.docx"))

    make_zip(
        root / "b.udf",
        {
            "content.xml": "<template><content><![CDATA[\nT.C. İŞÇİ\n]]></content></template>",
            "sign.sgn": b"sig",
        },
    )
    make_zip(root / "sheet.xlsx", {"xl/workbook.xml": "<x/>"})

    (root / "legacy.pdf").write_bytes(make_ole2("WordDocument", "1Table"))
    (root / "stub.pdf").write_bytes(b"<!doctype html>\n<html><body>index</body></html>")
    (root / "pic.jpg").write_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 64)
    (root / "note.txt").write_text("düz metin ğ\n", encoding="utf-8")
    (root / "Thumbs.db").write_bytes(make_ole2("Catalog"))
    (root / "broken.pdf").write_bytes(b"%PDF-1.4\ngarbage")
    return root


@pytest.fixture
def turkish_sentence() -> str:
    return _TURKISH_SENTENCE


@pytest.fixture
def one_text_file(tmp_path: Path) -> tuple[Path, Path, Path]:
    """(root, file, cache_dir) for a single extractable text file."""
    root = tmp_path / "corpus"
    root.mkdir()
    path = root / "a.txt"
    path.write_text("işçi ve işveren ile ilgili madde " * 5, encoding="utf-8")
    return root, path, tmp_path / "cache"


# --- Synthetic decision texts (task 04). Names, numbers and sentences are made up. ----------

_FILLER = "Örnek gerekçe cümlesi uydurma bir metindir. " * 25


def era2_text(
    court: str = "T.C\nYARGITAY\n9. HUKUK DAİRESİ",
    esas: str = "2011/12345",
    karar: str = "2012/6789",
    tarih: str = "07.06.2012",
    related: str = "4857 S. İşK/18-21\n1475 S. İşK/14",
    result: str = "Temyiz olunan kararın yukarıda yazılı sebepten BOZULMASINA, "
    "07.06.2012 gününde oybirliğiyle karar verildi.",
) -> str:
    """Journal header, İlgili Kanun, court block with inline labels, keywords, ÖZETİ, DAVA."""
    return (
        "Yargıtay Kararları\n117\n"
        f"İlgili Kanun / Madde\n{related}\n{court}\n"
        f"Esas No. {esas}\nKarar No. {karar}\nTarihi: {tarih}\n"
        "• ÖRNEK ANAHTAR BİR\n• ÖRNEK ANAHTAR İKİ\nSATIR KIRILIMI\n"
        "ÖZETİ: Örnek özet birinci satır\nikinci satır özet.\n"
        "DAVA: Davacı, örnek talebin kabulünü istemiştir.\n"
        "Yerel mahkeme, isteği reddetmiştir.\n"
        "Hüküm süresi içinde davacı vekili tarafından temyiz edilmiş olmakla dosya incelendi.\n"
        f"{_FILLER}\n\fYargıtay Kararları\n118\n{_FILLER}\n"
        f"SONUÇ: {result}"
    )


def era1_text() -> str:
    """Summary first: the court block with ESAS NO:/KARAR NO:/TARİHİ: follows the ÖZÜ."""
    return (
        "Yargıtay Kararları\n10\n"
        "İlgili Kanun/md:\n1475 s.İşK: 14\nBK:90\n"
        "• ÖRNEK ANAHTAR\nÖZÜ: Örnek özet satırı\nikinci satır.\n"
        "T.C.\nYARGITAY\n10. Hukuk Dairesi\n"
        "ESAS NO: 2003/111\nKARAR NO: 2003/222\nTARİHİ: 6.11.2003\n"
        "DAVA: Davacı, örnek talebin kabulünü istemiştir.\n"
        f"Yerel mahkeme, davayı reddetmiştir.\n{_FILLER}\n"
        "SONUÇ: Temyiz olunan hükmün yukarıda yazılı sebepten ONANMASINA, "
        "6.11.2003 gününde oybirliği ile karar verildi."
    )


def era3_text() -> str:
    """Two-column extraction: the labels in a row, then the values in the same order."""
    return (
        "YARGITAY\n22. HUKUK DAİRESİ\nEsas No.\nKarar No.\nTarihi:\n"
        "2014/90305\n2014/91998\n11.02.2014\n"
        "İlgili Kanun / Madde\n6356 S. STK/25,26\n"
        "• ÖRNEK ANAHTAR\nÖZETİ Örnek özet.\n"
        f"DAVA: Davacı, örnek talebin kabulünü istemiştir.\n{_FILLER}\n"
        "SONUÇ: Temyiz olunan kararın DÜZELTİLEREK ONANMASINA, "
        "11.02.2014 tarihinde oybirliğiyle karar verildi."
    )


def era4_text() -> str:
    """New Yargıtay template: roman-numbered sections."""
    return (
        "Yargıtay Kararları – Çalışma ve Toplum, 2024/1\n1\n"
        "İlgili Kanun / Madde\n4857 S. İşK/41,63,68\n"
        "T.C\nYARGITAY\n9. HUKUK DAİRESİ\n"
        "Esas No. 2023/99001\nKarar No. 2023/99002\nTarihi: 08.11.2023\n"
        "•ÖRNEK ANAHTAR\nÖZETİ: Örnek özet.\n"
        f"I. DAVA\nDavacı örnek talepte bulunmuştur.\nV. GEREKÇE\n{_FILLER}\n"
        "\fYargıtay Kararları – Çalışma ve Toplum, 2024/1\n2\n"
        "VI. KARAR\nTemyiz olunan kararın BOZULMASINA, 08.11.2023 tarihinde oy birliğiyle "
        "karar verildi."
    )


def foreign_text() -> str:
    return (
        "Alman Federal Mahkeme Kararları\n205\nHamm Eyalet İş Mahkemesi*\n"
        "Karar Tarihi : 29.06.2005\nSayısı : 14 Sa 469/05\n"
        "Örnek başlık cümlesi.\nÖzü:\nÖrnek özet.\n"
        f"Olay:\nÖrnek olay anlatımı. {_FILLER}\n"
    )
